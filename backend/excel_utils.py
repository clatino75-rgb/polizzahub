"""Excel import/export of policy anagrafiche using openpyxl."""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

FIELD_LABELS = [
    ("numero_polizza", "Numero Polizza"),
    ("compagnia_emissione", "Compagnia Emissione"),
    ("agenzia_emissione", "Agenzia Emissione"),
    ("tipo_polizza", "Tipo Polizza"),
    ("ramo_polizza", "Ramo Polizza"),
    ("contraente_nome", "Contraente"),
    ("contraente_cf_piva", "CF/P.IVA Contraente"),
    ("contraente_indirizzo", "Indirizzo Contraente"),
    ("proprietario_nome", "Proprietario"),
    ("proprietario_cf_piva", "CF/P.IVA Proprietario"),
    ("proprietario_indirizzo", "Indirizzo Proprietario"),
    ("data_effetto", "Data Effetto"),
    ("data_scadenza", "Data Scadenza"),
    ("frazionamento", "Frazionamento"),
    ("premio_netto_annuale", "Premio Netto Annuale"),
    ("premio_lordo_annuale", "Premio Lordo Annuale"),
    ("premio_netto_semestrale", "Premio Netto Semestrale"),
    ("premio_lordo_semestrale", "Premio Lordo Semestrale"),
    ("data_immatricolazione", "Data Immatricolazione"),
    ("data_voltura", "Data Voltura"),
    ("targa", "Targa"),
    ("note", "Note"),
]

KEYS = [k for k, _ in FIELD_LABELS]
LABEL_TO_KEY = {label.lower().strip(): key for key, label in FIELD_LABELS}
LABEL_TO_KEY.update({k: k for k in KEYS})


def policies_to_xlsx(policies: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anagrafiche"
    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, (_, label) in enumerate(FIELD_LABELS, start=1):
        c = ws.cell(row=1, column=col, value=label)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(16, len(label) + 2)
    for r, pol in enumerate(policies, start=2):
        for col, (key, _) in enumerate(FIELD_LABELS, start=1):
            ws.cell(row=r, column=col, value=str(pol.get(key, "") or ""))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_xlsx() -> bytes:
    return policies_to_xlsx([])


def xlsx_to_policies(file_bytes: bytes) -> list:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).lower().strip() if h is not None else "" for h in rows[0]]
    col_key = {}
    for idx, h in enumerate(header):
        if h in LABEL_TO_KEY:
            col_key[idx] = LABEL_TO_KEY[h]
    policies = []
    for row in rows[1:]:
        record = {}
        for idx, val in enumerate(row):
            key = col_key.get(idx)
            if key and val is not None:
                record[key] = str(val).strip()
        if any(record.values()):
            policies.append(record)
    return policies
